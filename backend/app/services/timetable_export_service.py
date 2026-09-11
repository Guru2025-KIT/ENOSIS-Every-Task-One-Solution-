"""
Timetable Export Service — Generates official PDF and Excel documents for Timetables.
Supports exporting by Division/Class, Faculty, and Room/Lab.
"""

import io
from typing import Dict, List, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
try:
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False



def generate_timetable_excel(
    college_name: str,
    department_name: str,
    academic_year: str,
    semester: str,
    view_title: str,
    days: List[str],
    time_slots: List[Dict[str, Any]],
    grid_data: Dict[str, List[str]] # key: f"{day}_{slot}", val: [subject, faculty, room/batch]
) -> bytes:
    """
    Generates a professionally styled Excel spreadsheet (.xlsx) for a timetable view.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Timetable"

    # Title Block
    ws.merge_cells("A1:G1")
    ws["A1"] = college_name or "ENOSIS Engineering Institute"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color="1E293B")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Department of {department_name or 'Computer Science'} | {academic_year or '2026-2027'} ({semester or 'Odd'} Semester)"
    ws["A2"].font = Font(name="Arial", size=11, italic=True, color="475569")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("A3:G3")
    ws["A3"] = f"Timetable View: {view_title}"
    ws["A3"].font = Font(name="Arial", size=12, bold=True, color="2563EB")
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")

    ws.append([]) # Empty row 4

    # Header Row (Row 5)
    headers = ["Time / Slot"] + [d[:3].upper() for d in days]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=5, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    break_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    lecture_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    lab_fill = PatternFill(start_color="F3E8FF", end_color="F3E8FF", fill_type="solid")

    for slot in time_slots:
        slot_num = slot.get("slot_number", 1)
        start_t = slot.get("start_time", "")
        end_t = slot.get("end_time", "")
        time_label = f"Slot {slot_num}\n({start_t} - {end_t})" if start_t else f"Slot {slot_num}"
        
        row = [time_label]
        is_break_slot = slot.get("is_break", False)

        for day in days:
            key = f"{day}_{slot_num}"
            cell_info = grid_data.get(key, ["Free", "", ""])
            subj = cell_info[0] if len(cell_info) > 0 else "Free"
            fac = cell_info[1] if len(cell_info) > 1 else ""
            extra = cell_info[2] if len(cell_info) > 2 else ""

            if is_break_slot or subj == "Break":
                row.append("BREAK")
            elif subj in ("Free", "-", ""):
                row.append("-")
            else:
                text = f"{subj}\n{fac}"
                if extra:
                    text += f"\n[{extra}]"
                row.append(text)

        ws.append(row)
        current_row = ws.max_row
        ws.row_dimensions[current_row].height = 45

        # Format cells
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = thin_border
            cell.font = Font(name="Arial", size=9)

            if col_num == 1:
                cell.font = Font(name="Arial", size=9, bold=True)
                cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            else:
                val = str(cell.value)
                if val == "BREAK":
                    cell.fill = break_fill
                    cell.font = Font(name="Arial", size=9, italic=True, color="64748B")
                elif "LAB" in val.upper() or "Batch" in val:
                    cell.fill = lab_fill
                elif val != "-":
                    cell.fill = lecture_fill

    # Set Column Widths
    ws.column_dimensions['A'].width = 18
    for col_letter in ['B', 'C', 'D', 'E', 'F', 'G']:
        ws.column_dimensions[col_letter].width = 22

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_timetable_pdf(
    college_name: str,
    department_name: str,
    academic_year: str,
    semester: str,
    view_title: str,
    days: List[str],
    time_slots: List[Dict[str, Any]],
    grid_data: Dict[str, List[str]]
) -> bytes:
    """
    Generates a clean PDF document for a timetable view using ReportLab.
    """
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError("The 'reportlab' package is not installed on the backend server. Please install reportlab to export PDFs.")

    buffer = io.BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=20,
        leftMargin=20,
        topMargin=20,
        bottomMargin=20
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        alignment=1, # Center
        textColor=colors.HexColor('#1E293B')
    )
    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        alignment=1,
        textColor=colors.HexColor('#475569')
    )
    view_style = ParagraphStyle(
        'ViewTitle',
        parent=styles['Heading2'],
        fontSize=12,
        leading=16,
        alignment=1,
        textColor=colors.HexColor('#2563EB')
    )

    cell_title_style = ParagraphStyle(
        'CellTitle',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        alignment=1,
        fontName='Helvetica-Bold'
    )
    cell_sub_style = ParagraphStyle(
        'CellSub',
        parent=styles['Normal'],
        fontSize=7,
        leading=8,
        alignment=1,
        textColor=colors.HexColor('#475569')
    )

    elements = [
        Paragraph(college_name or "ENOSIS Engineering Institute", title_style),
        Paragraph(f"Department of {department_name or 'Computer Science'} | {academic_year or '2026-2027'} ({semester or 'Odd'})", subtitle_style),
        Paragraph(f"Timetable View: {view_title}", view_style),
        Spacer(1, 10)
    ]

    # Build Table
    table_data = []
    header_row = ["Time / Slot"] + [d[:3].upper() for d in days]
    table_data.append(header_row)

    for slot in time_slots:
        slot_num = slot.get("slot_number", 1)
        start_t = slot.get("start_time", "")
        end_t = slot.get("end_time", "")
        time_text = f"Slot {slot_num}<br/>{start_t} - {end_t}" if start_t else f"Slot {slot_num}"
        
        row = [Paragraph(time_text, cell_title_style)]
        is_break_slot = slot.get("is_break", False)

        for day in days:
            key = f"{day}_{slot_num}"
            cell_info = grid_data.get(key, ["Free", "", ""])
            subj = cell_info[0] if len(cell_info) > 0 else "Free"
            fac = cell_info[1] if len(cell_info) > 1 else ""
            extra = cell_info[2] if len(cell_info) > 2 else ""

            if is_break_slot or subj == "Break":
                row.append(Paragraph("<b>BREAK</b>", cell_sub_style))
            elif subj in ("Free", "-", ""):
                row.append(Paragraph("-", cell_sub_style))
            else:
                html = f"<b>{subj}</b>"
                if fac:
                    html += f"<br/>{fac}"
                if extra:
                    html += f"<br/>[{extra}]"
                row.append(Paragraph(html, cell_sub_style))

        table_data.append(row)

    col_widths = [80] + [110] * len(days)
    t = Table(table_data, colWidths=col_widths, repeatRows=1)
    
    t_style = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0, 1), (0, -1), colors.HexColor('#F8FAFC')),
    ]
    t.setStyle(TableStyle(t_style))
    elements.append(t)

    doc.build(elements)
    return buffer.getvalue()
